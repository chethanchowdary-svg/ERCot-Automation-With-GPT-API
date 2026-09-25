"""Write the run outputs to one Excel workbook (one tab per output)."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import rules as R
from .engine import INVE, INVR, num

HDR_FILL = PatternFill("solid", start_color="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
WRAP_COLS = {"Analyst comments", "Proposed changes", "Detail", "Reason"}
CHANGED_FILL = PatternFill("solid", start_color="FFF2CC")   # light yellow = value changed


def _money(v):
    n = num(v) if v not in (None, "") else None
    return "" if n is None else (int(n) if float(n).is_integer() else round(n, 2))


def _sheet(wb, title, columns, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    changed_cells = []
    for n, r in enumerate(rows, start=2):
        for col in r.get("_changed", ()):
            if col in columns:
                changed_cells.append((n, columns.index(col) + 1))
        vals = []
        for col in columns:
            v = r.get(col, "")
            if col in (INVE, INVR):
                v = _money(v)
            vals.append(v)
        ws.append(vals)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
            if columns[c.column - 1] in WRAP_COLS:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(vertical="top")
    for rr, cc in changed_cells:
        ws.cell(rr, cc).fill = CHANGED_FILL
    for i, col in enumerate(columns, 1):
        w = (widths or {}).get(col) or (60 if col in WRAP_COLS else min(max(len(col) + 2, 12), 36))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write(res, path):
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Existing Update (Non-QC)", R.DB_COLUMNS, res.updates)
    _sheet(wb, "New Projects", R.NEW_COLUMNS, res.new)
    _sheet(wb, "Existing Update (QC)", R.DB_COLUMNS, res.qc)
    _sheet(wb, "Change Log & Flags",
           ["Type", "Flag", "Source sheet", "Interconnection ID", "Project ID", "Project name",
            "Lead QC", "Detail"], res.review, widths={"Flag": 48})
    if res.id_matches:
        _sheet(wb, "ID Matches (GPT)", list(res.id_matches[0].keys()), res.id_matches,
               widths={"Reason": 70, "ERCOT name": 32, "DB project name": 32})
    ws = wb.create_sheet("Run Summary")
    ws.append(["Metric", "Value"])
    for k, v in res.stats.items():
        ws.append([k, v])
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 42, 60
    wb.save(path)
